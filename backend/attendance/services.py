import csv
import openpyxl
import io
from io import TextIOWrapper
from django.http import FileResponse
from collections import defaultdict
from .models import Attendance, ClassBatch, Timetable, Notification


def generate_attendance_report(user, start_date, end_date, download_format=None):
    # 1. Filter based on Role
    if user.role == 'Student':
        qs = Attendance.objects.filter(student=user, date__range=[start_date, end_date])
        if download_format:
            raise PermissionError("Students are not allowed to download reports.")
    elif user.role == 'Teacher':
        qs = Attendance.objects.filter(class_batch__teacher=user, date__range=[start_date, end_date])
    else:  # Admin, Principal, HOD
        qs = Attendance.objects.filter(date__range=[start_date, end_date])

    if not qs.exists():
        return {"error": "No attendance data available for this date range."}

    # 2. Manual grouping and aggregation
    report_data = defaultdict(lambda: {'total_classes': 0, 'classes_attended': 0})
    
    for record in qs.select_related('student', 'class_batch__subject'):
        roll_no = record.student.roll_no if record.student else 'N/A'
        name = record.student.name if record.student else 'Unknown'
        subject = record.class_batch.subject.name if (record.class_batch and record.class_batch.subject) else 'Unknown'
        
        key = (roll_no, name, subject)
        report_data[key]['total_classes'] += 1
        if record.status == 'Present':
            report_data[key]['classes_attended'] += 1

    # 3. Compile report
    report_list = []
    for (roll_no, name, subject), stats in report_data.items():
        total = stats['total_classes']
        attended = stats['classes_attended']
        percentage = round((attended / total) * 100, 2) if total > 0 else 0.0
        
        report_list.append({
            'roll_number': roll_no,
            'name': name,
            'subject': subject,
            'total_classes': total,
            'classes_attended': attended,
            'attendance_percentage': percentage,
        })
        
    # 4. Return JSON for Frontend Charts (Recharts)
    if not download_format:
        return report_list

    # 5. Export FileResponse for Downloads
    if download_format == 'excel':
        buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        if ws is None:
            ws = wb.create_sheet()
        ws.title = 'Attendance Report'
        
        headers = ['Roll Number', 'Name', 'Subject', 'Total Classes', 'Classes Attended', 'Attendance (%)']
        ws.append(headers)
        
        for item in report_list:
            ws.append([
                item['roll_number'], item['name'], item['subject'], 
                item['total_classes'], item['classes_attended'], item['attendance_percentage']
            ])
            
        wb.save(buffer)
        buffer.seek(0)
        return FileResponse(buffer, as_attachment=True, filename='Attendance_Report.xlsx',
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    elif download_format == 'csv':
        buffer = io.BytesIO()
        text_buffer = io.StringIO()
        writer = csv.writer(text_buffer)
        
        headers = ['Roll Number', 'Name', 'Subject', 'Total Classes', 'Classes Attended', 'Attendance (%)']
        writer.writerow(headers)
        
        for item in report_list:
            writer.writerow([
                item['roll_number'], item['name'], item['subject'], 
                item['total_classes'], item['classes_attended'], item['attendance_percentage']
            ])
            
        buffer.write(text_buffer.getvalue().encode('utf-8'))
        buffer.seek(0)
        return FileResponse(buffer, as_attachment=True, filename='Attendance_Report.csv', content_type='text/csv')


def process_timetable_upload(file_obj):
    rows = []
    headers = []
    
    if file_obj.name.endswith('.csv'):
        csv_file = TextIOWrapper(file_obj.file, encoding='utf-8')
        reader = csv.DictReader(csv_file)
        headers = [h.strip() for h in reader.fieldnames] if reader.fieldnames else []
        for row in reader:
            rows.append({k.strip(): v for k, v in row.items() if k})
    elif file_obj.name.endswith(('.xls', '.xlsx')):
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        sheet = wb.active
        if sheet is None:
            raise ValueError("No active sheet in Excel file.")
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(col).strip() for col in row if col is not None]
            else:
                row_data = {}
                for j, col in enumerate(row):
                    if j < len(headers):
                        row_data[headers[j]] = col
                rows.append(row_data)
    else:
        raise ValueError("Invalid file format.")

    count = 0
    for row in rows:
        subject_id = row.get('Subject_ID')
        day = row.get('Day')
        start_time = row.get('Start_Time')
        end_time = row.get('End_Time')
        
        if not all([subject_id, day, start_time, end_time]):
            continue
            
        try:
            batch = ClassBatch.objects.get(id=subject_id)
            timetable, created = Timetable.objects.update_or_create(
                class_batch=batch,
                day_of_week=day,
                start_time=start_time,
                defaults={'end_time': end_time}
            )
            count += 1

            # Notification Logic
            subject_name = batch.subject.name if batch.subject else "Subject"
            msg = f"Your timetable for {subject_name} has been updated."

            # 1. Notify Teacher
            if batch.teacher:
                Notification.objects.create(user=batch.teacher, message=msg)

            # 2. Notify Students
            enrollments = batch.enrollments.select_related('student')
            notifications = [
                Notification(user=enrollment.student, message=msg)
                for enrollment in enrollments if enrollment.student
            ]
            if notifications:
                Notification.objects.bulk_create(notifications)
        except ClassBatch.DoesNotExist:
            continue
    return count